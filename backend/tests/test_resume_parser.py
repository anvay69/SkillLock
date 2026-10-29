from io import BytesIO
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.schemas.profile import (
    Certification,
    Education,
    Experience,
    ExtractedProfile,
    Project,
)


def make_messy_layout_docx() -> bytes:
    document = Document()
    document.add_paragraph("Ananya Rao | ananya@example.com | +91 98765 43210")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Technical Skills: Python, Go, Kafka"
    table.cell(0, 1).text = "Professional Experience: Staff Engineer at Example Corp"
    table.cell(1, 0).text = "Academic Background: B.Tech Computer Science, 2022"
    table.cell(1, 1).text = "Selected Projects: Event Processor | Kafka, Go"
    document.add_paragraph("Personal Details: Home address must not be extracted")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class ResumeParserApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.client = TestClient(app)
        self.extracted = ExtractedProfile(
            skills=["Python", "Go", "Kafka"],
            projects=[
                Project(
                    title="Event Processor",
                    description="Processes events from a partitioned stream.",
                    tech_stack="Go, Kafka",
                )
            ],
            experience=[
                Experience(
                    job_title="Staff Engineer",
                    company="Example Corp",
                    start_date="2021",
                )
            ],
            education=[
                Education(
                    institution="Example University",
                    degree="B.Tech",
                    field_of_study="Computer Science",
                    graduation_year="2022",
                )
            ],
            certifications=[Certification(name="Cloud Engineer")],
        )

    def test_extracts_professional_details_from_messy_table_layout(self) -> None:
        with (
            patch("app.features.resume_parser.routes.get_session_user"),
            patch(
                "app.features.resume.service.extract_resume_profile",
                return_value=self.extracted,
            ) as extract,
        ):
            response = self.client.post(
                "/api/resume-parser/parse",
                files={
                    "resume": (
                        "messy-resume.docx",
                        make_messy_layout_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )

        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["skills"], ["Python", "Go", "Kafka"])
        self.assertEqual(body["projects"][0]["title"], "Event Processor")
        self.assertEqual(body["experience"][0]["job_title"], "Staff Engineer")
        self.assertEqual(body["education"][0]["degree"], "B.Tech")
        self.assertEqual(body["certifications"][0]["name"], "Cloud Engineer")
        self.assertEqual(set(body), {"skills", "projects", "experience", "education", "certifications"})

        sections = extract.call_args.args[0]
        self.assertEqual(sections["skills"], ["Python, Go, Kafka"])
        self.assertEqual(sections["experience"], ["Staff Engineer at Example Corp"])
        self.assertEqual(sections["education"], ["B.Tech Computer Science, 2022"])
        self.assertEqual(sections["projects"], ["Event Processor | Kafka, Go"])
        self.assertNotIn("Ananya Rao", str(sections))
        self.assertNotIn("ananya@example.com", str(sections))
        self.assertNotIn("Home address", str(sections))
        self.assertEqual(extract.call_args.args[1].name, "")
        self.assertEqual(extract.call_args.args[1].email, "")

    def test_requires_authentication_and_pdf_or_docx(self) -> None:
        with TestClient(app) as anonymous:
            unauthorized = anonymous.post(
                "/api/resume-parser/parse",
                files={"resume": ("resume.docx", make_messy_layout_docx())},
            )
        self.assertEqual(unauthorized.status_code, 401)

        with patch("app.features.resume_parser.routes.get_session_user"):
            unsupported = self.client.post(
                "/api/resume-parser/parse",
                files={"resume": ("resume.txt", b"Resume details")},
            )
        self.assertEqual(unsupported.status_code, 400)
        self.assertIn("PDF or DOCX", unsupported.json()["detail"])


if __name__ == "__main__":
    unittest.main()
